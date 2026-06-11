"""Assemble the output workbook from extracted rows, following the fixed layout.

Input is a mapping {sheet_name: list[dict]} where each dict is keyed by the
sheet's column labels (from schema.py). Every sheet in the schema is written even
if it has no rows. Per-sheet layout: row 1 title, row 2 note, row 4 headers,
row 5+ data.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .schema import (DATA_START_ROW, HEADER_ROW, NOTE_ROW, SHEETS, TITLE_ROW,
                     SheetSpec)

_TITLE_FONT = Font(bold=True, size=12)
_NOTE_FONT = Font(italic=True, size=9, color="555555")
_HEADER_FONT = Font(bold=True)


def _write_sheet(ws, spec: SheetSpec, rows: list[dict]) -> None:
    ncols = len(spec.columns)
    ws.cell(row=TITLE_ROW, column=1, value=spec.title).font = _TITLE_FONT
    if ncols > 1:
        ws.merge_cells(start_row=TITLE_ROW, start_column=1,
                       end_row=TITLE_ROW, end_column=ncols)
    if spec.note:
        c = ws.cell(row=NOTE_ROW, column=1, value=spec.note)
        c.font = _NOTE_FONT
        if ncols > 1:
            ws.merge_cells(start_row=NOTE_ROW, start_column=1,
                           end_row=NOTE_ROW, end_column=ncols)

    for j, col in enumerate(spec.columns, start=1):
        ws.cell(row=HEADER_ROW, column=j, value=col).font = _HEADER_FONT

    for i, row in enumerate(rows):
        r = DATA_START_ROW + i
        for j, col in enumerate(spec.columns, start=1):
            ws.cell(row=r, column=j, value=row.get(col))

    for j, col in enumerate(spec.columns, start=1):
        width = min(max(len(str(col)) + 2, 10), 40)
        ws.column_dimensions[get_column_letter(j)].width = width


def write_workbook(out_path: Path, sheet_rows: dict[str, list[dict]]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for spec in SHEETS:
        ws = wb.create_sheet(title=spec.name)
        _write_sheet(ws, spec, sheet_rows.get(spec.name, []))
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
