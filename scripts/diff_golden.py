"""Cross-check a produced UNIVERSAL workbook against the GEICO golden.

Reshapes the universal sheets (06_Vehicle_Symbols, 05_Territory_Definitions,
04_Rating_Factors) back into the GEICO-shaped sheets and compares. The golden is
NOT authoritative (it was produced by another model run) -- a mismatch means
"verify against the PDF page", not "we are wrong". See project memory.
"""

import sys
from pathlib import Path

import openpyxl

from rate_filing import geico_view
from rate_filing.schema import DATA_START_ROW, SHEET_BY_NAME


def _rows_as_dicts(ws, columns: list[str]) -> list[dict]:
    out = []
    for r in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
        if not any(c is not None and str(c).strip() for c in r):
            continue
        out.append({col: ("" if r[i] is None else r[i])
                    for i, col in enumerate(columns) if i < len(r)})
    return out


def _golden_rows(ws):
    out = []
    for r in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
        if any(c is not None and str(c).strip() for c in r):
            out.append(tuple("" if c is None else str(c).strip() for c in r))
    return out


def main(produced: str, golden: str) -> int:
    pw = openpyxl.load_workbook(produced, read_only=True)
    gw = openpyxl.load_workbook(golden, read_only=True)

    veh = _rows_as_dicts(pw["06_Vehicle_Symbols"], SHEET_BY_NAME["06_Vehicle_Symbols"].columns)
    terr = _rows_as_dicts(pw["05_Territory_Definitions"], SHEET_BY_NAME["05_Territory_Definitions"].columns)
    fac = _rows_as_dicts(pw["04_Rating_Factors"], SHEET_BY_NAME["04_Rating_Factors"].columns)

    views = {
        "19_MBI_Codes": geico_view.mbi_view(veh),
        "20_VLR_Symbols": geico_view.vlr_view(veh),
        "18_Territory_Definitions": geico_view.territory_view(terr),
        "14_Class_Factors": geico_view.class_factors_view(fac),
    }

    print(f"{'golden sheet':<28} {'mine':>7} {'golden':>7}  status")
    for name, mine in views.items():
        gold = _golden_rows(gw[name]) if name in gw.sheetnames else []
        status = ""
        if name == "18_Territory_Definitions":
            mk = {str(r["ZIP Code"]): str(r["Rating Territory"]) for r in mine}
            gk = {r[4]: r[5] for r in gold}
            wrong = sum(1 for z in mk.keys() & gk.keys() if mk[z] != gk[z])
            missing = len(gk.keys() - mk.keys())
            status = (f"0 wrong values; {missing} not matched"
                      if wrong == 0 else f"{wrong} wrong values")
        print(f"{name:<28} {len(mine):>7} {len(gold):>7}  {status}")
    return 0


if __name__ == "__main__":
    p = sys.argv[1] if len(sys.argv) > 1 else "output/geico.xlsx"
    g = sys.argv[2] if len(sys.argv) > 2 else "fixtures/GEICO_NJ_PPA_Rates_Rules_RAG_v4.xlsx"
    raise SystemExit(main(p, g))
