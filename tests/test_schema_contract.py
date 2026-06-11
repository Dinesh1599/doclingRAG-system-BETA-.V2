"""The output workbook must always present every sheet and column in order."""

from pathlib import Path

import openpyxl

from rate_filing.excel_writer import write_workbook
from rate_filing.schema import HEADER_ROW, SHEETS, SHEET_NAMES


def test_all_sheets_and_columns_present(tmp_path: Path):
    out = write_workbook(tmp_path / "wb.xlsx", {})
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == SHEET_NAMES
    for spec in SHEETS:
        ws = wb[spec.name]
        header = [c.value for c in ws[HEADER_ROW][: len(spec.columns)]]
        assert header == spec.columns, f"{spec.name} header mismatch: {header}"
