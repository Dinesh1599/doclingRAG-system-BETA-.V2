"""Regression tests for GEICO extraction, via the universal builders + the
GEICO-view reshaper, cross-checked against the (non-authoritative) golden.

Skipped unless the GEICO PDF, its docling cache, and the golden are present.
NOTE: the golden was produced by another model run, so a mismatch means "go
check the PDF", not necessarily "we are wrong" (see project memory).
"""

import json
from pathlib import Path

import openpyxl
import pytest

from rate_filing import geico_view
from rate_filing import sheet_builders as sb
from rate_filing.discover import discover
from rate_filing.quarantine import Quarantine

PDF = Path("input/geico.pdf")
CACHE = Path("input/geico.pdf.docling.json")
GOLDEN = Path("fixtures/GEICO_NJ_PPA_Rates_Rules_RAG_v4.xlsx")
pytestmark = pytest.mark.skipif(
    not (PDF.exists() and CACHE.exists() and GOLDEN.exists()),
    reason="GEICO fixtures / docling cache absent")

MD = {"Company": "GEICO Indemnity Company", "NAIC": "22055",
      "State": "NJ", "Line": "Private Passenger Auto"}


def _doc():
    return json.loads(CACHE.read_text())


def test_territory_values_match_golden_by_zip():
    d = discover(PDF, log=lambda *a, **k: None)
    rows = sb.build_territory(_doc(), d, MD, Quarantine(), log=lambda *a, **k: None)
    view = geico_view.territory_view(rows)
    gw = openpyxl.load_workbook(GOLDEN, read_only=True)
    gold = {str(r[4]): str(r[5]) for r in gw["18_Territory_Definitions"]
            .iter_rows(min_row=5, values_only=True) if r[4]}
    mine = {r["ZIP Code"]: str(r["Rating Territory"]) for r in view}
    wrong = [z for z in mine.keys() & gold.keys() if mine[z] != gold[z]]
    assert wrong == [], f"wrong territory values: {wrong[:5]}"


def test_class_factors_cover_all_golden_values():
    d = discover(PDF, log=lambda *a, **k: None)
    factors = sb.build_rating_factors(PDF, _doc(), d, MD, Quarantine(),
                                      log=lambda *a, **k: None)
    view = geico_view.class_factors_view(factors)
    gw = openpyxl.load_workbook(GOLDEN, read_only=True)
    gold = {}
    for r in gw["14_Class_Factors"].iter_rows(min_row=5, values_only=True):
        if r[4]:
            gold[(r[4], str(r[5]))] = tuple(float(x) for x in r[6:10])
    mine = {(r["Coverage"], str(r["Age of Driver"])):
            (r["Single Male"], r["Single Female"], r["Married Male"], r["Married Female"])
            for r in view}
    for key, gvals in gold.items():
        assert key in mine, f"missing class-factor row {key}"
        mvals = tuple(float(str(v).replace(",", "")) for v in mine[key])
        assert mvals == pytest.approx(gvals), f"{key}: {mvals} != {gvals}"
